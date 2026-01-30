"""
Generador d'informes tècnics de reg.

Genera informes en formats:
- PDF (reportlab)
- DOCX (python-docx)
- Markdown
"""

from pathlib import Path
from typing import Dict, List, Optional
from datetime import datetime, date
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm, mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    Image as RLImage, PageBreak
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY

from ..cultius import TipusCultiu, NOMS_CULTIUS
from ..config import EstatHidrologic


class ReportGenerator:
    """Generador d'informes tècnics de reg."""

    def __init__(self, output_dir: Optional[Path] = None):
        """
        Inicialitza el generador.

        Args:
            output_dir: Directori de sortida per als informes
        """
        self.output_dir = output_dir or Path("output/reports")
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _format_data(self, data: date) -> str:
        """Formata una data en català."""
        mesos = [
            'gener', 'febrer', 'març', 'abril', 'maig', 'juny',
            'juliol', 'agost', 'setembre', 'octubre', 'novembre', 'desembre'
        ]
        return f"{data.day} de {mesos[data.month-1]} de {data.year}"

    def generar_markdown(self, dades: Dict) -> str:
        """
        Genera un informe en format Markdown.

        Args:
            dades: Diccionari amb totes les dades

        Returns:
            Text en format Markdown
        """
        cultiu_nom = NOMS_CULTIUS.get(
            TipusCultiu(dades['cultiu']['tipus']),
            dades['cultiu']['tipus']
        )

        md = []

        # Capçalera
        md.append("# INFORME TÈCNIC DE REG")
        md.append(f"## Canals d'Urgell - {dades['municipi']}")
        md.append(f"**Data:** {self._format_data(date.today())}")
        md.append(f"**Cultiu:** {cultiu_nom}")
        md.append("")

        # 1. Estat del sistema hidrològic
        md.append("---")
        md.append("## 1️⃣ Estat del Sistema Hidrològic")
        md.append("")

        hidro = dades['estat_hidrologic']
        for emb in hidro.get('embassaments', []):
            md.append(f"### Embassament de {emb['nom']}")
            md.append(f"- **Volum:** {emb['volum_hm3']:.1f} hm³ / {emb['capacitat_hm3']:.1f} hm³")
            md.append(f"- **Percentatge:** {emb['percentatge']:.1f}%")
            md.append("")

        sistema = hidro.get('sistema', {})
        md.append(f"**Estat global:** {sistema.get('percentatge_global', 0):.1f}%")
        md.append(f"**Classificació:** {sistema.get('estat', 'desconegut').upper()}")
        md.append(f"**Diagnòstic:** {sistema.get('estat_descripcio', '')}")
        md.append("")

        sequera = hidro.get('sequera', {})
        if sequera.get('mode_actiu'):
            md.append(f"⚠️ **MODE SEQUERA ACTIU** - Reducció del {sequera.get('reduccio_percentatge', 0):.0f}%")
            md.append("")

        # 2. Anàlisi agrometeorològica
        md.append("---")
        md.append("## 2️⃣ Anàlisi Agrometeorològica")
        md.append("")

        agrometeo = dades.get('agrometeo', {})
        md.append(f"| Paràmetre | Valor |")
        md.append(f"|-----------|-------|")
        md.append(f"| Temperatura actual | {agrometeo.get('temperatura_actual', 0):.1f} °C |")
        md.append(f"| Humitat relativa | {agrometeo.get('humitat', 0):.0f}% |")
        md.append(f"| Velocitat del vent | {agrometeo.get('vent_ms', 0):.1f} m/s |")
        md.append(f"| ETo diària | {agrometeo.get('eto_mm', 0):.2f} mm |")
        md.append(f"| Radiació solar | {agrometeo.get('radiacio_solar', 0):.1f} MJ/m²/dia |")
        md.append("")

        pluja = dades.get('pluja', {})
        md.append("### Precipitació")
        md.append(f"- Última hora: {pluja.get('ultima_hora_mm', 0):.1f} mm")
        md.append(f"- Prevista 24h: {pluja.get('prevista_24h_mm', 0):.1f} mm")
        md.append(f"- Prevista 48h: {pluja.get('prevista_48h_mm', 0):.1f} mm")
        md.append("")

        # 3. Necessitat hídrica del cultiu
        md.append("---")
        md.append("## 3️⃣ Necessitat Hídrica del Cultiu")
        md.append("")

        cultiu_info = dades.get('cultiu', {})
        md.append(f"**Fase fenològica:** {cultiu_info.get('fase_fenologica', '')}")
        md.append(f"**Coeficient Kc:** {cultiu_info.get('kc_actual', 0):.2f}")
        md.append(f"**Dies des de sembra:** {cultiu_info.get('dies_desde_sembra', 0)}")
        md.append("")

        necessitat = dades.get('necessitat_avui', {})
        md.append("### Balanç Hídric (FAO-56)")
        md.append(f"| Component | Valor (mm) |")
        md.append(f"|-----------|------------|")
        md.append(f"| ETc (ETo × Kc) | {necessitat.get('etc_mm', 0):.2f} |")
        md.append(f"| Pluja efectiva | {necessitat.get('pluja_efectiva_mm', 0):.2f} |")
        md.append(f"| Dèficit hídric | {necessitat.get('deficit_mm', 0):.2f} |")
        md.append(f"| **Dosi neta** | **{necessitat.get('dosi_neta_mm', 0):.2f}** |")
        md.append(f"| **Dosi bruta** | **{necessitat.get('dosi_bruta_mm', 0):.2f}** |")
        md.append("")

        # 4. Recomanació operativa
        md.append("---")
        md.append("## 4️⃣ Recomanació Operativa")
        md.append("")

        if necessitat.get('regar'):
            md.append(f"### ✅ REGAR AVUI: {necessitat.get('dosi_bruta_mm', 0):.1f} mm")
            md.append(f"**Equivalent:** {necessitat.get('dosi_m3_ha', 0):.0f} m³/ha")
        else:
            md.append("### ❌ NO CAL REGAR AVUI")

        md.append(f"**Prioritat:** {necessitat.get('prioritat', 'baixa').upper()}")
        md.append(f"**Motiu:** {necessitat.get('motiu', '')}")
        md.append("")

        if necessitat.get('mode_sequera'):
            md.append(f"⚠️ Reg de supervivència aplicat (factor {necessitat.get('factor_sequera', 1):.2f})")
            md.append("")

        # 5. Justificació tècnica
        md.append("---")
        md.append("## 5️⃣ Justificació Tècnica")
        md.append("")
        md.append("### Metodologia FAO-56 (Penman-Monteith)")
        md.append("")
        md.append("El càlcul de les necessitats de reg segueix l'estàndard internacional:")
        md.append("")
        md.append("```")
        md.append("ETc = ETo × Kc")
        md.append(f"ETc = {agrometeo.get('eto_mm', 0):.2f} × {cultiu_info.get('kc_actual', 0):.2f} = {necessitat.get('etc_mm', 0):.2f} mm")
        md.append("")
        md.append("Necessitat de Reg = ETc - Pluja Efectiva")
        md.append(f"Necessitat = {necessitat.get('etc_mm', 0):.2f} - {necessitat.get('pluja_efectiva_mm', 0):.2f} = {necessitat.get('deficit_mm', 0):.2f} mm")
        md.append("```")
        md.append("")

        sistema_reg = dades.get('sistema_reg', {})
        md.append(f"**Sistema de reg:** {sistema_reg.get('tipus', 'aspersio')}")
        md.append(f"**Eficiència:** {sistema_reg.get('eficiencia', 0.75)*100:.0f}%")
        md.append(f"**Dosi bruta = Dosi neta / Eficiència = {necessitat.get('dosi_neta_mm', 0):.2f} / {sistema_reg.get('eficiencia', 0.75):.2f} = {necessitat.get('dosi_bruta_mm', 0):.2f} mm**")
        md.append("")

        # Pla setmanal
        md.append("---")
        md.append("## 📅 Pla de Reg Setmanal")
        md.append("")

        pla = dades.get('pla_setmanal', {})
        dies = pla.get('dies', [])

        if dies:
            md.append("| Data | ETo | ETc | Pluja | Reg | Prioritat |")
            md.append("|------|-----|-----|-------|-----|-----------|")
            for dia in dies:
                data_str = dia.get('data', '')[:10]
                regar = "✅" if dia.get('regar') else "—"
                md.append(
                    f"| {data_str} | {dia.get('eto_mm', 0):.1f} | {dia.get('etc_mm', 0):.1f} | "
                    f"{dia.get('pluja_mm', 0):.1f} | {regar} {dia.get('dosi_bruta_mm', 0):.1f} | {dia.get('prioritat', '')} |"
                )
            md.append("")
            md.append(f"**Total reg setmana:** {pla.get('total_reg_mm', 0):.1f} mm ({pla.get('total_reg_m3_ha', 0):.0f} m³/ha)")
            md.append("")

        # Alertes
        alertes = pla.get('alertes', []) + dades.get('alertes', {}).get('onada_calor', [])
        if alertes or dades.get('alertes', {}).get('onada_calor'):
            md.append("---")
            md.append("## ⚠️ Alertes")
            md.append("")
            for alerta in alertes:
                md.append(f"- {alerta}")
            if dades.get('alertes', {}).get('onada_calor'):
                md.append(f"- 🌡️ Onada de calor prevista: {dades['alertes'].get('dies_calor_extrema', 0)} dies >35°C")
            md.append("")

        # Peu
        md.append("---")
        md.append("*Informe generat automàticament per l'Agent de Reg dels Canals d'Urgell*")
        md.append(f"*{datetime.now().strftime('%Y-%m-%d %H:%M')}*")

        return "\n".join(md)

    def generar_docx(
        self,
        dades: Dict,
        grafics: Optional[List[Path]] = None,
        filename: Optional[str] = None
    ) -> Path:
        """
        Genera un informe en format DOCX.

        Args:
            dades: Diccionari amb totes les dades
            grafics: Llista de paths a gràfics PNG
            filename: Nom del fitxer de sortida

        Returns:
            Path del fitxer generat
        """
        if filename is None:
            filename = f"informe_reg_{date.today().isoformat()}.docx"

        filepath = self.output_dir / filename

        doc = Document()

        # Estils
        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)

        # Títol
        titulo = doc.add_heading('INFORME TÈCNIC DE REG', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitulo = doc.add_paragraph()
        subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitulo.add_run(f"Canals d'Urgell - {dades['municipi']}")
        run.bold = True
        run.font.size = Pt(14)

        doc.add_paragraph(f"Data: {self._format_data(date.today())}")

        cultiu_nom = NOMS_CULTIUS.get(
            TipusCultiu(dades['cultiu']['tipus']),
            dades['cultiu']['tipus']
        )
        doc.add_paragraph(f"Cultiu: {cultiu_nom}")

        # 1. Estat hidrològic
        doc.add_heading('1. Estat del Sistema Hidrològic', level=1)

        hidro = dades['estat_hidrologic']
        for emb in hidro.get('embassaments', []):
            p = doc.add_paragraph()
            p.add_run(f"Embassament de {emb['nom']}: ").bold = True
            p.add_run(f"{emb['volum_hm3']:.1f} hm³ / {emb['capacitat_hm3']:.1f} hm³ ({emb['percentatge']:.1f}%)")

        sistema = hidro.get('sistema', {})
        p = doc.add_paragraph()
        p.add_run("Estat global: ").bold = True
        p.add_run(f"{sistema.get('estat', 'desconegut').upper()} - {sistema.get('estat_descripcio', '')}")

        sequera = hidro.get('sequera', {})
        if sequera.get('mode_actiu'):
            p = doc.add_paragraph()
            run = p.add_run(f"⚠️ MODE SEQUERA ACTIU - Reducció del {sequera.get('reduccio_percentatge', 0):.0f}%")
            run.bold = True

        # 2. Anàlisi agrometeorològica
        doc.add_heading('2. Anàlisi Agrometeorològica', level=1)

        agrometeo = dades.get('agrometeo', {})
        table = doc.add_table(rows=5, cols=2)
        table.style = 'Table Grid'

        cells = table.rows[0].cells
        cells[0].text = 'Temperatura actual'
        cells[1].text = f"{agrometeo.get('temperatura_actual', 0):.1f} °C"

        cells = table.rows[1].cells
        cells[0].text = 'Humitat relativa'
        cells[1].text = f"{agrometeo.get('humitat', 0):.0f}%"

        cells = table.rows[2].cells
        cells[0].text = 'Velocitat del vent'
        cells[1].text = f"{agrometeo.get('vent_ms', 0):.1f} m/s"

        cells = table.rows[3].cells
        cells[0].text = 'ETo diària'
        cells[1].text = f"{agrometeo.get('eto_mm', 0):.2f} mm"

        cells = table.rows[4].cells
        cells[0].text = 'Radiació solar'
        cells[1].text = f"{agrometeo.get('radiacio_solar', 0):.1f} MJ/m²/dia"

        # 3. Necessitat hídrica
        doc.add_heading('3. Necessitat Hídrica del Cultiu', level=1)

        necessitat = dades.get('necessitat_avui', {})
        cultiu_info = dades.get('cultiu', {})

        doc.add_paragraph(f"Fase fenològica: {cultiu_info.get('fase_fenologica', '')}")
        doc.add_paragraph(f"Coeficient Kc: {cultiu_info.get('kc_actual', 0):.2f}")

        table = doc.add_table(rows=5, cols=2)
        table.style = 'Table Grid'

        rows_data = [
            ('ETc (ETo × Kc)', f"{necessitat.get('etc_mm', 0):.2f} mm"),
            ('Pluja efectiva', f"{necessitat.get('pluja_efectiva_mm', 0):.2f} mm"),
            ('Dèficit hídric', f"{necessitat.get('deficit_mm', 0):.2f} mm"),
            ('Dosi neta', f"{necessitat.get('dosi_neta_mm', 0):.2f} mm"),
            ('Dosi bruta', f"{necessitat.get('dosi_bruta_mm', 0):.2f} mm"),
        ]

        for i, (label, value) in enumerate(rows_data):
            cells = table.rows[i].cells
            cells[0].text = label
            cells[1].text = value

        # 4. Recomanació operativa
        doc.add_heading('4. Recomanació Operativa', level=1)

        if necessitat.get('regar'):
            p = doc.add_paragraph()
            run = p.add_run(f"✅ REGAR AVUI: {necessitat.get('dosi_bruta_mm', 0):.1f} mm ({necessitat.get('dosi_m3_ha', 0):.0f} m³/ha)")
            run.bold = True
            run.font.size = Pt(14)
        else:
            p = doc.add_paragraph()
            run = p.add_run("❌ NO CAL REGAR AVUI")
            run.bold = True
            run.font.size = Pt(14)

        doc.add_paragraph(f"Prioritat: {necessitat.get('prioritat', 'baixa').upper()}")
        doc.add_paragraph(f"Motiu: {necessitat.get('motiu', '')}")

        # 5. Justificació
        doc.add_heading('5. Justificació Tècnica (FAO-56)', level=1)

        doc.add_paragraph(
            f"ETc = ETo × Kc = {agrometeo.get('eto_mm', 0):.2f} × {cultiu_info.get('kc_actual', 0):.2f} = {necessitat.get('etc_mm', 0):.2f} mm"
        )
        doc.add_paragraph(
            f"Necessitat = ETc - Pluja efectiva = {necessitat.get('etc_mm', 0):.2f} - {necessitat.get('pluja_efectiva_mm', 0):.2f} = {necessitat.get('deficit_mm', 0):.2f} mm"
        )

        # Gràfics
        if grafics:
            doc.add_page_break()
            doc.add_heading('Gràfics', level=1)

            for grafic in grafics:
                if grafic.exists():
                    doc.add_picture(str(grafic), width=Inches(6))
                    doc.add_paragraph()

        # Guardar
        doc.save(filepath)
        return filepath

    def generar_pdf(
        self,
        dades: Dict,
        grafics: Optional[List[Path]] = None,
        filename: Optional[str] = None
    ) -> Path:
        """
        Genera un informe en format PDF.

        Args:
            dades: Diccionari amb totes les dades
            grafics: Llista de paths a gràfics PNG
            filename: Nom del fitxer de sortida

        Returns:
            Path del fitxer generat
        """
        if filename is None:
            filename = f"informe_reg_{date.today().isoformat()}.pdf"

        filepath = self.output_dir / filename

        doc = SimpleDocTemplate(
            str(filepath),
            pagesize=A4,
            rightMargin=2*cm,
            leftMargin=2*cm,
            topMargin=2*cm,
            bottomMargin=2*cm
        )

        styles = getSampleStyleSheet()

        # Estils personalitzats
        styles.add(ParagraphStyle(
            name='TitolPrincipal',
            parent=styles['Heading1'],
            fontSize=18,
            alignment=TA_CENTER,
            spaceAfter=20
        ))

        styles.add(ParagraphStyle(
            name='Subtitol',
            parent=styles['Normal'],
            fontSize=14,
            alignment=TA_CENTER,
            spaceAfter=10
        ))

        styles.add(ParagraphStyle(
            name='Seccio',
            parent=styles['Heading2'],
            fontSize=14,
            spaceBefore=15,
            spaceAfter=10
        ))

        story = []

        # Títol
        story.append(Paragraph("INFORME TÈCNIC DE REG", styles['TitolPrincipal']))
        story.append(Paragraph(f"Canals d'Urgell - {dades['municipi']}", styles['Subtitol']))
        story.append(Paragraph(f"Data: {self._format_data(date.today())}", styles['Normal']))

        cultiu_nom = NOMS_CULTIUS.get(
            TipusCultiu(dades['cultiu']['tipus']),
            dades['cultiu']['tipus']
        )
        story.append(Paragraph(f"Cultiu: {cultiu_nom}", styles['Normal']))
        story.append(Spacer(1, 20))

        # 1. Estat hidrològic
        story.append(Paragraph("1. Estat del Sistema Hidrològic", styles['Seccio']))

        hidro = dades['estat_hidrologic']
        data_emb = [['Embassament', 'Volum (hm³)', 'Capacitat (hm³)', '%']]
        for emb in hidro.get('embassaments', []):
            data_emb.append([
                emb['nom'],
                f"{emb['volum_hm3']:.1f}",
                f"{emb['capacitat_hm3']:.1f}",
                f"{emb['percentatge']:.1f}%"
            ])

        if len(data_emb) > 1:
            t = Table(data_emb, colWidths=[4*cm, 3*cm, 3.5*cm, 2*cm])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            story.append(t)

        sistema = hidro.get('sistema', {})
        story.append(Spacer(1, 10))
        story.append(Paragraph(
            f"<b>Estat global:</b> {sistema.get('estat', 'desconegut').upper()} - {sistema.get('estat_descripcio', '')}",
            styles['Normal']
        ))

        # 2. Agrometeorologia
        story.append(Paragraph("2. Anàlisi Agrometeorològica", styles['Seccio']))

        agrometeo = dades.get('agrometeo', {})
        data_meteo = [
            ['Paràmetre', 'Valor'],
            ['Temperatura actual', f"{agrometeo.get('temperatura_actual', 0):.1f} °C"],
            ['Humitat relativa', f"{agrometeo.get('humitat', 0):.0f}%"],
            ['Velocitat del vent', f"{agrometeo.get('vent_ms', 0):.1f} m/s"],
            ['ETo diària', f"{agrometeo.get('eto_mm', 0):.2f} mm"],
            ['Radiació solar', f"{agrometeo.get('radiacio_solar', 0):.1f} MJ/m²/dia"],
        ]

        t = Table(data_meteo, colWidths=[5*cm, 5*cm])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightblue),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        story.append(t)

        # 3. Necessitat hídrica
        story.append(Paragraph("3. Necessitat Hídrica del Cultiu", styles['Seccio']))

        necessitat = dades.get('necessitat_avui', {})
        cultiu_info = dades.get('cultiu', {})

        story.append(Paragraph(
            f"Fase fenològica: {cultiu_info.get('fase_fenologica', '')} | Kc: {cultiu_info.get('kc_actual', 0):.2f}",
            styles['Normal']
        ))

        data_balanc = [
            ['Component', 'Valor (mm)'],
            ['ETc (ETo × Kc)', f"{necessitat.get('etc_mm', 0):.2f}"],
            ['Pluja efectiva', f"{necessitat.get('pluja_efectiva_mm', 0):.2f}"],
            ['Dèficit hídric', f"{necessitat.get('deficit_mm', 0):.2f}"],
            ['Dosi neta', f"{necessitat.get('dosi_neta_mm', 0):.2f}"],
            ['Dosi bruta', f"{necessitat.get('dosi_bruta_mm', 0):.2f}"],
        ]

        t = Table(data_balanc, colWidths=[5*cm, 5*cm])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgreen),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('BACKGROUND', (-1, -1), (-1, -1), colors.yellow),
            ('FONTNAME', (-1, -1), (-1, -1), 'Helvetica-Bold'),
        ]))
        story.append(t)

        # 4. Recomanació
        story.append(Paragraph("4. Recomanació Operativa", styles['Seccio']))

        if necessitat.get('regar'):
            story.append(Paragraph(
                f"<b>✅ REGAR AVUI: {necessitat.get('dosi_bruta_mm', 0):.1f} mm ({necessitat.get('dosi_m3_ha', 0):.0f} m³/ha)</b>",
                styles['Normal']
            ))
        else:
            story.append(Paragraph("<b>❌ NO CAL REGAR AVUI</b>", styles['Normal']))

        story.append(Paragraph(f"Prioritat: {necessitat.get('prioritat', 'baixa').upper()}", styles['Normal']))
        story.append(Paragraph(f"Motiu: {necessitat.get('motiu', '')}", styles['Normal']))

        # 5. Justificació
        story.append(Paragraph("5. Justificació Tècnica (FAO-56)", styles['Seccio']))

        story.append(Paragraph(
            f"ETc = ETo × Kc = {agrometeo.get('eto_mm', 0):.2f} × {cultiu_info.get('kc_actual', 0):.2f} = {necessitat.get('etc_mm', 0):.2f} mm",
            styles['Normal']
        ))

        # Gràfics
        if grafics:
            story.append(PageBreak())
            story.append(Paragraph("Gràfics", styles['Seccio']))

            for grafic in grafics:
                if grafic.exists():
                    img = RLImage(str(grafic), width=16*cm, height=10*cm)
                    story.append(img)
                    story.append(Spacer(1, 20))

        # Generar PDF
        doc.build(story)
        return filepath
